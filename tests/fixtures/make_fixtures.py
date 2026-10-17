# -*- coding: utf-8 -*-
"""テスト用のファイルを生成する（生成物はリポジトリに含めるので、通常は実行不要）

    python make_fixtures.py

必要なもの: openpyxl、LibreOffice（soffice）。
LibreOffice を通すのは、Excel と同じく共有文字列テーブル（sharedStrings.xml）や
スタイル付きの日付を持つ xlsx にするため（openpyxl は文字列をインラインで書く）。
"""
import datetime
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile

import openpyxl
from openpyxl.comments import Comment

HERE = os.path.dirname(os.path.abspath(__file__))
GREP = os.path.join(HERE, "grep")
PDF = os.path.join(HERE, "pdf")
REVIEW = os.path.join(HERE, "review")
REPO = os.path.dirname(os.path.dirname(HERE))


def soffice_convert(src, fmt, outdir):
    subprocess.run(["soffice", "--headless", "--convert-to", fmt, "--outdir", outdir, src],
                   check=True, capture_output=True, timeout=300)


def rewrite_zip(path, files):
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for n, b in files.items():
            z.writestr(n, b)


def read_zip(path):
    with zipfile.ZipFile(path) as z:
        return {n: z.read(n) for n in z.namelist()}


DRAWING = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">
<xdr:twoCellAnchor><xdr:from><xdr:col>6</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>2</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from><xdr:to><xdr:col>8</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>5</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to>
<xdr:sp><xdr:nvSpPr><xdr:cNvPr id="2" name="テキスト ボックス 1"/><xdr:cNvSpPr txBox="1"/></xdr:nvSpPr><xdr:spPr/><xdr:txBody><a:bodyPr/><a:p><a:r><a:t>注意：</a:t></a:r><a:r><a:t>TODO 図形内</a:t></a:r></a:p><a:p><a:r><a:t>2行目</a:t></a:r></a:p></xdr:txBody></xdr:sp><xdr:clientData/></xdr:twoCellAnchor>
<mc:AlternateContent><mc:Choice Requires="a14"><xdr:twoCellAnchor><xdr:from><xdr:col>1</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>9</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from><xdr:to><xdr:col>2</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>10</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to><xdr:sp><xdr:nvSpPr><xdr:cNvPr id="3" name="Button 1"/><xdr:cNvSpPr/></xdr:nvSpPr><xdr:spPr/><xdr:txBody><a:bodyPr/><a:p><a:r><a:t>TODO実行</a:t></a:r></a:p></xdr:txBody></xdr:sp><xdr:clientData/></xdr:twoCellAnchor></mc:Choice>
<mc:Fallback><xdr:twoCellAnchor><xdr:from><xdr:col>1</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>9</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from><xdr:to><xdr:col>2</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>10</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to><xdr:sp><xdr:nvSpPr><xdr:cNvPr id="3" name="Button 1"/><xdr:cNvSpPr/></xdr:nvSpPr><xdr:spPr/><xdr:txBody><a:bodyPr/><a:p><a:r><a:t>TODO実行</a:t></a:r></a:p></xdr:txBody></xdr:sp><xdr:clientData/></xdr:twoCellAnchor></mc:Fallback></mc:AlternateContent>
</xdr:wsDr>"""


def make_grep_fixtures(tmp):
    """Search-Excel.ps1 用。期待値は tests/Search-Excel.Tests.ps1 を参照"""
    shutil.rmtree(GREP, ignore_errors=True)
    os.makedirs(os.path.join(GREP, "sub"))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "顧客一覧"
    ws["A1"] = "顧客ID"; ws["B1"] = "会社名"; ws["C1"] = "登録日"; ws["D1"] = "時刻"; ws["E1"] = "金額"
    ws["A2"] = "C-001"; ws["B2"] = "(株)山田商事"
    ws["C2"] = datetime.date(2026, 10, 16); ws["C2"].number_format = "yyyy/m/d"
    ws["D2"] = datetime.time(13, 0); ws["D2"].number_format = "h:mm"
    ws["E2"] = 0.1 + 0.2
    ws["A3"] = "C-002"; ws["B3"] = "ＡＢＣ株式会社"
    ws["C3"] = datetime.datetime(2026, 1, 2, 9, 30); ws["C3"].number_format = '[$-411]ggge"年"m"月"d"日" h:mm'
    ws["E3"] = "=SUM(E2,1)"
    ws["B4"] = "複数行\nの文字列 TODO"
    ws["B5"] = True
    ws["E5"] = "=VLOOKUP(A2,顧客一覧!A:B,2,FALSE)"
    ws["A2"].comment = Comment("要確認：住所が古い", "山田")
    ws["F1"] = 12345678901
    hidden = wb.create_sheet("作業用")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "隠しTODO"
    big = wb.create_sheet("Sheet3")
    for r in range(1, 2001):
        big.cell(row=r, column=1, value="行{0}".format(r))
        big.cell(row=r, column=2, value=r * 1.5)
        big.cell(row=r, column=3, value="=B{0}*2".format(r))
    big["A2000"] = "最終行TODO"
    base = os.path.join(tmp, "base.xlsx")
    wb.save(base)

    soffice_convert(base, "xlsx", os.path.join(tmp, "lo"))
    files = read_zip(os.path.join(tmp, "lo", "base.xlsx"))

    # ふりがな（rPh）付きの共有文字列にする
    sst = files["xl/sharedStrings.xml"].decode()
    sst = sst.replace('<t xml:space="preserve">会社名</t>',
                      '<r><t>会社</t></r><r><t>名</t></r><rPh sb="0" eb="2"><t>カイシャメイフリガナ</t></rPh><phoneticPr fontId="1"/>', 1)
    assert "カイシャ" in sst
    files["xl/sharedStrings.xml"] = sst.encode()

    # 図形（mc:Fallback に同じ図形が重複して入る形）を追加する
    files["xl/drawings/drawing99.xml"] = DRAWING.encode()
    rels_name = "xl/worksheets/_rels/sheet1.xml.rels"
    rels = files[rels_name].decode().replace("\n</Relationships>", "</Relationships>")
    rels = rels.replace("</Relationships>", '<Relationship Id="rIdD99" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing" Target="/xl/drawings/drawing99.xml"/></Relationships>')
    files[rels_name] = rels.encode()
    sheet = files["xl/worksheets/sheet1.xml"].decode()
    sheet = re.sub(r"(<legacyDrawing)", r'<drawing xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="rIdD99"/>\1', sheet)
    assert "rIdD99" in sheet
    files["xl/worksheets/sheet1.xml"] = sheet.encode()
    ct = files["[Content_Types].xml"].decode().replace("</Types>", '<Override PartName="/xl/drawings/drawing99.xml" ContentType="application/vnd.openxmlformats-officedocument.drawing+xml"/></Types>')
    files["[Content_Types].xml"] = ct.encode()

    sample = os.path.join(GREP, "sample.xlsx")
    rewrite_zip(sample, files)

    shutil.copy(sample, os.path.join(GREP, "~$sample.xlsx"))                 # Excel の一時ファイル（対象外）
    shutil.copy(sample, os.path.join(GREP, "sub", "disguised.xls"))          # 中身は xlsx の .xls
    with open(os.path.join(GREP, "sub", "encrypted.xlsx"), "wb") as f:      # 暗号化 xlsx と同じ先頭（OLE）
        f.write(bytes.fromhex("D0CF11E0A1B11AE1") + b"\0" * 600)
    with open(os.path.join(GREP, "sub", "broken.xlsx"), "wb") as f:         # 壊れた ZIP
        f.write(b"PK\x03\x04garbage")

    # インライン文字列・r 属性なし・_x000D_ エスケープ
    inline = os.path.join(GREP, "sub", "inline.xlsx")
    wb2 = openpyxl.Workbook(); wb2.active.title = "inline"; wb2.save(inline)
    files = read_zip(inline)
    files["xl/worksheets/sheet1.xml"] = (
        '<?xml version="1.0" encoding="UTF-8"?><worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>'
        '<row><c t="inlineStr"><is><t>インライン_x000D_TODO</t></is></c><c t="inlineStr"><is><t>二列目</t></is></c></row>'
        '<row><c><v>3</v></c><c t="inlineStr"><is><t>B2のTODO</t></is></c></row>'
        '</sheetData></worksheet>').encode()
    rewrite_zip(inline, files)

    # シート名に ' を含む・制御文字（_x0001_）を含むセル
    wb3 = openpyxl.Workbook(); ws3 = wb3.active; ws3.title = "O'Brien 表 1"
    ws3["C5"] = 'TODO <&> "quote" _x0001_ctrl'
    wb3.save(os.path.join(GREP, "sub", "quote.xlsx"))

    # 本物の .xls（Excel が無い環境では読めないことの確認用）
    wb4 = openpyxl.Workbook(); wb4.active["A1"] = "旧形式TODO"
    legacy_src = os.path.join(tmp, "legacy.xlsx"); wb4.save(legacy_src)
    soffice_convert(legacy_src, "xls", os.path.join(GREP, "sub"))


def make_pdf(path, sizes):
    """sizes の各要素（幅, 高さ: pt）を 1 ページとする最小限の PDF を書く"""
    objs = []
    kids = " ".join("{0} 0 R".format(3 + i) for i in range(len(sizes)))
    objs.append("<< /Type /Catalog /Pages 2 0 R >>")
    objs.append("<< /Type /Pages /Kids [{0}] /Count {1} >>".format(kids, len(sizes)))
    for w, h in sizes:
        objs.append("<< /Type /Page /Parent 2 0 R /MediaBox [0 0 {0} {1}] >>".format(w, h))
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, o in enumerate(objs, start=1):
        offsets.append(len(out))
        out += "{0} 0 obj\n{1}\nendobj\n".format(i, o).encode()
    xref = len(out)
    out += "xref\n0 {0}\n0000000000 65535 f \n".format(len(objs) + 1).encode()
    for off in offsets:
        out += "{0:010d} 00000 n \n".format(off).encode()
    out += "trailer\n<< /Size {0} /Root 1 0 R >>\nstartxref\n{1}\n%%EOF\n".format(len(objs) + 1, xref).encode()
    with open(path, "wb") as f:
        f.write(out)


def make_pdf_fixtures():
    """Export-ExcelToPdf.ps1 -CountOnly 用"""
    shutil.rmtree(PDF, ignore_errors=True)
    os.makedirs(PDF)
    a4 = (595, 842)
    make_pdf(os.path.join(PDF, "a4x3.pdf"), [a4, a4, a4])
    make_pdf(os.path.join(PDF, "mixed.pdf"), [a4, (842, 595), (100, 100)])   # 3 ページ目だけ極端に小さい
    with open(os.path.join(PDF, "broken.pdf"), "wb") as f:
        f.write(b"%PDF-1.4\nthis is not a pdf\n")


def make_review_fixtures():
    """Invoke-DesignReview.ps1 用。同梱の観点表に誤りを入れた観点表を作る"""
    shutil.rmtree(REVIEW, ignore_errors=True)
    os.makedirs(REVIEW)
    src = os.path.join(REPO, "tools", "review-checklists", "詳細設計レビュー観点.xlsx")
    wb = openpyxl.load_workbook(src)
    ws = wb.active
    # 列: A観点ID B分類 C観点 D判定基準 E重要度 F対象設計書 G参照設計書 H有効
    ws["A3"] = "s01"                                   # 5 行目の S01 と大文字小文字違いの重複
    ws["C4"] = None                                    # 観点が空欄
    ws["D6"] = None                                    # 判定基準が空欄
    ws["D7"] = "初期値が適切に設定されている（例: 必要に応じて「なし」）"  # あいまい（例の中は対象外）
    ws["E8"] = "最高"                                  # 重要度
    ws["H9"] = "はい"                                  # 有効の値
    ws["A10"] = "S 08"                                 # 空白入りの ID
    ws["F11"] = "*存在しない設計書*"                  # 当たらないパターン
    ws["G12"] = "*存在しない参照*"
    ws["D13"] = "A | B のどちらかが記載"               # 区切り文字
    ws["A14"] = "S01"; ws["H14"] = "×"                 # 無効な行の重複はエラーにしない
    wb.save(os.path.join(REVIEW, "broken-checklist.xlsx"))


DELIVERABLE = os.path.join(HERE, "deliverable")


def add_parts(path, parts, content_types):
    """xlsx に任意のパーツを足す（parts: 名前 → 中身、content_types: 追加する Override / Default 要素）"""
    files = read_zip(path)
    for name, body in parts.items():
        if name in files and name.endswith(".rels") and body.startswith("+"):
            files[name] = files[name].decode().replace("</Relationships>", body[1:] + "</Relationships>").encode()
        else:
            files[name] = body.encode() if isinstance(body, str) else body
    ct = files["[Content_Types].xml"].decode()
    files["[Content_Types].xml"] = ct.replace("</Types>", "".join(content_types) + "</Types>").encode()
    rewrite_zip(path, files)


def make_deliverable_fixtures(tmp):
    """Test-ExcelDeliverable.ps1 用。good.xlsx は指摘なし、bad.xlsx は各チェックに 1 つ以上引っかかる"""
    shutil.rmtree(DELIVERABLE, ignore_errors=True)
    os.makedirs(DELIVERABLE)

    # ---- 指摘なし: A1 選択・倍率 100%・先頭シート表示・見出し固定のシート（スクロール領域の左上を選択）
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "表紙"; ws["A1"] = "提出用"
    ws2 = wb.create_sheet("一覧"); ws2["A1"] = "No"; ws2["A2"] = 1
    ws2.freeze_panes = "A2"
    ws2.sheet_view.selection[0].activeCell = "A2"; ws2.sheet_view.selection[0].sqref = "A2"
    wb.properties.creator = None
    wb.properties.lastModifiedBy = None
    wb.save(os.path.join(DELIVERABLE, "good.xlsx"))

    # ---- 指摘あり
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "表紙"
    ws["A1"] = "提出用"
    ws["C5"] = "TODO: 数値を確認"
    ws["B2"] = "#REF!"                                    # エラー値
    ws["D1"] = "メモ付き"
    ws["D1"].comment = Comment("社内向けの注意", "山田")
    ws.sheet_view.selection[0].activeCell = "C5"; ws.sheet_view.selection[0].sqref = "C5"
    ws.sheet_view.topLeftCell = "B3"
    ws.sheet_view.zoomScale = 85
    ws.row_dimensions[8].hidden = True
    ws["A10"] = "資料"
    ws["A10"].hyperlink = "file:///C:/Users/yamada/Desktop/memo.xlsx"
    ws2 = wb.create_sheet("詳細")
    ws2["A1"] = "詳細"
    ws2.sheet_view.view = "pageBreakPreview"
    ws3 = wb.create_sheet("作業用"); ws3.sheet_state = "hidden"
    ws4 = wb.create_sheet("隠し"); ws4.sheet_state = "veryHidden"
    wb.active = 1                                          # 2 枚目が表示された状態
    ws.sheet_view.tabSelected = True                       # 1 枚目も選択 → 作業グループ
    ws2.sheet_view.tabSelected = True
    from openpyxl.workbook.defined_name import DefinedName
    wb.defined_names["壊れた名前"] = DefinedName("壊れた名前", attr_text="#REF!")
    wb.properties.creator = "山田 太郎"
    wb.properties.lastModifiedBy = "佐藤 花子"
    bad = os.path.join(DELIVERABLE, "bad.xlsx")
    wb.save(bad)
    # 外部リンク（他ブックの参照）と会社名のプロパティを足す
    files = read_zip(bad)
    app = files["docProps/app.xml"].decode()
    app = app.replace("</Properties>", "<Company>株式会社サンプル</Company></Properties>")
    files["docProps/app.xml"] = app.encode()
    rewrite_zip(bad, files)
    add_parts(bad, {
        "xl/externalLinks/externalLink1.xml":
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><externalLink xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><externalBook r:id="rId1"><sheetNames><sheetName val="Sheet1"/></sheetNames></externalBook></externalLink>',
        "xl/externalLinks/_rels/externalLink1.xml.rels":
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/externalLinkPath" Target="file:///\\\\fileserver\\share\\原価表.xlsx" TargetMode="External"/></Relationships>',
        "xl/_rels/workbook.xml.rels":
            '+<Relationship Id="rIdX1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/externalLink" Target="externalLinks/externalLink1.xml"/>',
    }, ['<Override PartName="/xl/externalLinks/externalLink1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.externalLink+xml"/>'])

    # ---- マクロ付き（中身は空の vbaProject.bin）
    wb = openpyxl.Workbook(); wb.active["A1"] = "マクロ"
    wb.properties.creator = None
    macro = os.path.join(DELIVERABLE, "macro.xlsm")
    wb.save(macro)
    add_parts(macro, {"xl/vbaProject.bin": b"\0" * 16},
              ['<Default Extension="bin" ContentType="application/vnd.ms-office.vbaProject"/>'])

    # ---- 読めないファイル
    shutil.copy(os.path.join(GREP, "sub", "encrypted.xlsx"), os.path.join(DELIVERABLE, "encrypted.xlsx"))
    shutil.copy(os.path.join(GREP, "sub", "legacy.xls"), os.path.join(DELIVERABLE, "legacy.xls"))


VIEW = os.path.join(HERE, "view")


def make_view_fixtures():
    """Reset-ExcelView.ps1 用。見出し（1 行目・A 列）を固定したまま、下へスクロールして D60 を選択した状態"""
    shutil.rmtree(VIEW, ignore_errors=True)
    os.makedirs(VIEW)
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "一覧"
    for r in range(1, 101):
        ws.cell(row=r, column=1, value=r)
        ws.cell(row=r, column=4, value="値{0}".format(r))
    ws.freeze_panes = "B2"
    wb.properties.creator = None
    path = os.path.join(VIEW, "frozen.xlsx")
    wb.save(path)
    files = read_zip(path)
    sheet = files["xl/worksheets/sheet1.xml"].decode()
    sheet = sheet.replace('topLeftCell="B2" activePane="bottomRight"', 'topLeftCell="B50" activePane="bottomRight"')
    sheet = sheet.replace('<selection pane="bottomRight" activeCell="A1" sqref="A1"/>', '<selection pane="bottomRight" activeCell="D60" sqref="D60"/>')
    sheet = sheet.replace('<selection pane="topRight"/>', '<selection pane="topRight" activeCell="D1" sqref="D1"/>')
    assert 'B50' in sheet and 'D60' in sheet
    files["xl/worksheets/sheet1.xml"] = sheet.encode()
    rewrite_zip(path, files)


COMPARE = os.path.join(HERE, "compare")


def make_compare_fixtures():
    """Compare-Excel.ps1 用。v1 → v2 で、桁数の変更・行の挿入と削除・数式の変更・コメントの追加・
    シートの追加と削除・ファイルの追加と削除を起こす"""
    shutil.rmtree(COMPARE, ignore_errors=True)
    os.makedirs(os.path.join(COMPARE, "v1")); os.makedirs(os.path.join(COMPARE, "v2"))

    def book(path, rows, extra_sheet=None, formula="=SUM(C2:C4)", comment=None):
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = "項目定義"
        for r, row in enumerate(rows, start=1):
            for c, v in enumerate(row, start=1):
                ws.cell(row=r, column=c, value=v)
        ws.cell(row=len(rows) + 1, column=3, value=formula)
        if comment:
            ws["B2"].comment = Comment(comment, "x")
        if extra_sheet:
            wb.create_sheet(extra_sheet)["A1"] = "{0}の内容".format(extra_sheet)
        wb.properties.creator = None
        wb.save(path)

    old_rows = [["No", "項目名", "桁数"], [1, "受注番号", 12], [2, "顧客コード", 8], [3, "顧客名", 100], [4, "備考", 500]]
    new_rows = [["No", "項目名", "桁数"], [1, "受注番号", 12], [2, "顧客コード", 10], [2.5, "顧客区分", 1], [3, "顧客名", 100]]
    v1 = os.path.join(COMPARE, "v1"); v2 = os.path.join(COMPARE, "v2")
    book(os.path.join(v1, "項目定義.xlsx"), old_rows, extra_sheet="旧メモ")
    book(os.path.join(v2, "項目定義.xlsx"), new_rows, extra_sheet="変更履歴", formula="=SUM(C2:C5)", comment="桁数を見直し")
    book(os.path.join(v1, "同じ.xlsx"), old_rows); book(os.path.join(v2, "同じ.xlsx"), old_rows)
    book(os.path.join(v1, "削除された.xlsx"), old_rows)
    book(os.path.join(v2, "追加された.xlsx"), new_rows)


def make_compare_format_fixtures(tmp):
    """Compare-Excel.ps1 -CompareFormat 用（compare/format）。old → new で、先頭に 1 行挿入したうえで
    太字・塗りつぶし・部分書式・取り消し線を変える。期待値は tests/Compare-Excel.Tests.ps1 を参照。
    openpyxl は文字列をセルに直接（inlineStr）書くので、LibreOffice で共有文字列にした版も sst に作る"""
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    from openpyxl.styles import Color, Font, PatternFill

    out = os.path.join(COMPARE, "format")
    os.makedirs(out, exist_ok=True)
    red = InlineFont(color="FFFF0000")
    rows = [["No", "項目名", "説明"], [1, "受注番号", "受注を一意に識別する番号"], [2, "受注日", "受注を受け付けた日付"],
            [3, "顧客名", None], [4, "顧客区分", None], [5, "備考", "自由記述"], [6, "更新者", "最終更新者"], [7, "更新日時", "最終更新日時"]]

    def book(path, new):
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = "画面項目"
        data = [list(r) for r in rows]
        if new:
            data.insert(1, [0, "受注ID", "新規に追加した項目"])
        for r, row in enumerate(data, start=1):
            for c, v in enumerate(row, start=1):
                ws.cell(row=r, column=c, value=v)
        off = 1 if new else 0
        for c in range(1, 4):
            ws.cell(row=1, column=c).font = Font(bold=True)
            ws.cell(row=1, column=c).fill = PatternFill("solid", fgColor="FFD9D9D9")
        if new:
            ws.cell(row=2 + off, column=2).font = Font(bold=True)                   # 太字にした
        else:
            for c in range(1, 4):
                ws.cell(row=3, column=c).fill = PatternFill("solid", fgColor="FFFFFF00")   # 塗りつぶしをやめた
        # 部分書式: 旧は「正式名称」だけ赤、新は書式なし（文字は同じ）
        ws.cell(row=4 + off, column=3).value = "顧客の正式名称" if new else CellRichText("顧客の", TextBlock(red, "正式名称"))
        # 文字も部分書式も変えた（部分書式の差分は出さず、値の差分だけになる）
        ws.cell(row=5 + off, column=3).value = (CellRichText(TextBlock(red, "個人・法人・団体"), "の区分") if new
                                                else CellRichText(TextBlock(red, "個人・法人"), "の区分"))
        if not new:
            ws.cell(row=6, column=2).font = Font(color=Color(theme=1))              # テキスト1 は自動と同じ扱い
        if new:
            for r in (7 + off, 8 + off):
                ws.cell(row=r, column=3).font = Font(strike=True)                 # 取り消し線を付けた
        wb.properties.creator = None
        wb.save(path)

    book(os.path.join(out, "old.xlsx"), False)
    book(os.path.join(out, "new.xlsx"), True)
    sst = os.path.join(out, "sst")
    os.makedirs(sst, exist_ok=True)
    for name in ("old.xlsx", "new.xlsx"):
        soffice_convert(os.path.join(out, name), "xlsx", tmp)
        shutil.move(os.path.join(tmp, name), os.path.join(sst, name))


REPLACE = os.path.join(HERE, "replace")


def make_replace_fixtures(tmp):
    """Replace-Excel.ps1 用（replace）。期待値は tests/Replace-Excel.Tests.ps1 を参照。
    shared.xlsx: LibreOffice で共有文字列にしたブック。表紙と本文で同じ文字列「受注管理システム」を共有し、ふりがなを付ける
    table.xlsx: openpyxl のブック（文字列はセルに直接）。テーブルの見出し・数式・コメントを持つ"""
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    from openpyxl.worksheet.table import Table

    shutil.rmtree(REPLACE, ignore_errors=True)
    os.makedirs(REPLACE)

    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "表紙"
    ws["A1"] = "受注管理システム 基本設計書"; ws["A2"] = "受注管理システム"; ws["A3"] = "第1.0版"
    body = wb.create_sheet("本文")
    body["A1"] = "受注管理システム"
    body["B1"] = CellRichText("顧客の", TextBlock(InlineFont(color="FFFF0000"), "正式名称"))
    body["C1"] = "変更しない"
    src = os.path.join(tmp, "shared.xlsx")
    wb.properties.creator = None
    wb.save(src)
    out = os.path.join(tmp, "out")
    os.makedirs(out, exist_ok=True)
    soffice_convert(src, "xlsx", out)
    path = os.path.join(REPLACE, "shared.xlsx")
    files = read_zip(os.path.join(out, "shared.xlsx"))
    sst = files["xl/sharedStrings.xml"].decode()
    target = '<si><t xml:space="preserve">受注管理システム</t></si>'
    assert target in sst, sst[:500]
    # ふりがな（Excel で日本語を入力すると付く）
    sst = sst.replace(target, '<si><t xml:space="preserve">受注管理システム</t><rPh sb="0" eb="2"><t>ジュチュウ</t></rPh><phoneticPr fontId="1"/></si>')
    files["xl/sharedStrings.xml"] = sst.encode()
    rewrite_zip(path, files)

    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "一覧"
    ws.append(["TODO列", "数量", "合計"])
    ws.append(["TODO値", 1, "=SUM(B2:B3)"])
    ws.append(["その他", 2, None])
    ws.add_table(Table(displayName="T_LIST", ref="A1:C3"))
    ws["A2"].comment = Comment("TODO: 確認", "x")
    wb.properties.creator = None
    wb.save(os.path.join(REPLACE, "table.xlsx"))


TERMS = os.path.join(HERE, "terms")


def make_terms_fixtures():
    """Test-ExcelTerms.ps1 と Replace-Excel.ps1 -ListPath 用（terms）。期待値は tests/Test-ExcelTerms.Tests.ps1 を参照"""
    shutil.rmtree(TERMS, ignore_errors=True)
    os.makedirs(TERMS)

    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "機能"
    ws["A1"] = "ユーザー管理"; ws["A2"] = "ユーザ登録"; ws["A3"] = "サーバー構成"; ws["A4"] = "ＩＤで検索"
    ws["A5"] = "ﾕｰｻﾞｰ一覧"; ws["A6"] = "インターフェース仕様"; ws["A7"] = "Id は対象外（大文字・小文字の違い）"
    wb.properties.creator = None
    wb.save(os.path.join(TERMS, "a.xlsx"))

    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "画面"
    ws["A1"] = "ユーザーIDを入力"; ws["A2"] = "サーバに接続"; ws["A3"] = "インタフェース"
    ws["A4"] = "変更出来る"; ws["A5"] = "削除できる"; ws["A6"] = "Id"
    wb.properties.creator = None
    wb.save(os.path.join(TERMS, "b.xlsx"))

    with open(os.path.join(TERMS, "用語集.csv"), "w", encoding="utf-8-sig", newline="") as f:
        f.write("正しい表記,揺れ,備考\r\nユーザー,ユーザ,外来語の長音\r\nできる,出来る/出来ます,ひらがなで書く\r\n")


PRINT = os.path.join(HERE, "print")


def make_print_fixtures():
    """Test-ExcelDeliverable.ps1 の印刷・版数のチェック用（print）。期待値は tests/Test-ExcelDeliverable.Tests.ps1 を参照"""
    from openpyxl.worksheet.properties import PageSetupProperties

    shutil.rmtree(PRINT, ignore_errors=True)
    os.makedirs(PRINT)

    # wide.xlsx: 横に広い表（A4 縦で約 5 ページ）、1 ページに収める設定のシート、印刷範囲の外にデータがあるシート
    wb = openpyxl.Workbook()
    wide = wb.active; wide.title = "一覧"
    fit = wb.create_sheet("収める")
    area = wb.create_sheet("範囲")
    for ws in (wide, fit):
        for c in range(1, 21):
            ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = 20
            ws.cell(row=1, column=c, value="列{0}".format(c))
        ws.page_setup.paperSize = 9
        ws.page_setup.orientation = "portrait"
    fit.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    fit.page_setup.fitToWidth = 1
    fit.page_setup.fitToHeight = 0
    wide.oddHeader.center.text = "TODO 社内用"
    wide.oddFooter.left.text = "&Z&F"
    fit.oddFooter.center.text = "&P / &N"
    for r in range(1, 11):
        for c in range(1, 4):
            area.cell(row=r, column=c, value=r * c)
    area["E20"] = "メモ"
    area.print_area = "A1:C10"
    area.page_setup.paperSize = 9
    area.page_setup.orientation = "landscape"
    wb.properties.creator = None
    wb.save(os.path.join(PRINT, "wide.xlsx"))

    def versioned(path, cover, history_header, history):
        wb = openpyxl.Workbook()
        ws = wb.active; ws.title = "表紙"
        for ref, value in cover.items():
            ws[ref] = value
        h = wb.create_sheet("改訂履歴")
        h["A1"] = "改訂履歴"
        for c, v in enumerate(history_header, start=1):
            h.cell(row=3, column=c, value=v)
        for r, row in enumerate(history, start=4):
            for c, v in enumerate(row, start=1):
                h.cell(row=r, column=c, value=v)
        wb.properties.creator = None
        wb.save(path)

    # 表紙の版数・更新日と、改訂履歴の最後の行、ファイル名の版数がそろっていない。版数が前の行より小さい行もある
    versioned(os.path.join(PRINT, "仕様書_v1.2.xlsx"),
              {"B2": "仕様書", "B5": "版数", "D5": "1.0", "B6": "最終更新日", "D6": datetime.date(2026, 10, 1)},
              ["版数", "改訂日", "内容"],
              [["1.0", datetime.date(2026, 10, 1), "新規作成"], ["1.2", datetime.date(2026, 10, 5), "見直し"], ["1.1", datetime.date(2026, 10, 12), "修正"]])
    # 改訂履歴のシートの上部（帳票のヘッダ）にも「版数」「更新日」がある。表の見出しは 3 行目
    path = os.path.join(PRINT, "帳票ヘッダ付き.xlsx")
    versioned(path, {"B5": "版数", "D5": "1.1", "B6": "更新日", "D6": datetime.date(2026, 10, 12)},
              ["版数", "改訂日", "内容"], [["1.0", datetime.date(2026, 10, 1), "新規作成"], ["1.1", datetime.date(2026, 10, 12), "修正"]])
    wb = openpyxl.load_workbook(path)
    h = wb["改訂履歴"]
    h["E1"] = "版数"; h["F1"] = "1.1"; h["G1"] = "更新日"; h["H1"] = datetime.date(2026, 10, 12)
    wb.properties.creator = None
    wb.save(path)
    # そろっている（表紙は「第1.1版」の形、日付は文字、改訂履歴の版数は Ver.1.1 の形）
    versioned(os.path.join(PRINT, "仕様書_第1.1版.xlsx"),
              {"B2": "仕様書", "B3": "第1.1版", "B6": "更新日", "D6": "2026年10月12日"},
              ["版", "日付", "内容"],
              [["Ver.1.0", "2026/10/1", "新規作成"], ["Ver.1.1", "2026/10/12", "修正"]])


SIZE = os.path.join(HERE, "size")


def make_size_fixtures():
    """Measure-DesignSize.ps1 用（size）。期待値は tests/Measure-DesignSize.Tests.ps1 を参照"""
    shutil.rmtree(SIZE, ignore_errors=True)
    os.makedirs(SIZE)

    def header(ws, labels):
        # 帳票のヘッダ（見出しと値の組）を 1 行目に置く
        c = 1
        for label, value in labels:
            ws.cell(row=1, column=c, value=label); ws.cell(row=1, column=c + 1, value=value)
            c += 3

    def table(ws, top, title, headers, rows):
        ws.cell(row=top, column=1, value="■ " + title)
        for i, h in enumerate(headers, start=1):
            ws.cell(row=top + 1, column=i, value=h)
        for r, row in enumerate(rows, start=top + 2):
            for i, v in enumerate(row, start=1):
                ws.cell(row=r, column=i, value=v)

    def save(wb, name):
        wb.properties.creator = None
        wb.save(os.path.join(SIZE, name))

    # 画面: 項目 3・イベント 3（EV02 は登録あり＝外部入力、EV03 は参照だけ＝外部照会、EV01 はテーブルなし）
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "画面項目定義"
    header(ws, [("画面ID", "S01"), ("画面名", "ログイン")])
    table(ws, 3, "画面項目定義", ["No", "項目名", "項目ID"], [[1, "ユーザーID", "userId"], [2, "パスワード", "password"], [3, "ログインボタン", "login"]])
    ev = wb.create_sheet("イベント定義")
    header(ev, [("画面ID", "S01"), ("画面名", "ログイン")])
    table(ev, 3, "イベント定義", ["イベントID", "イベント名", "使用テーブル"],
          [["EV01", "初期表示", None], ["EV02", "ログイン", "M_USER(R)、T_LOGIN_LOG(C)"], ["EV03", "ユーザー検索", "M_USER（R）"]])
    table(ev, 10, "補足", ["メモ"], [["次の節は数えない"]])
    save(wb, "画面設計書_S01_ログイン.xlsx")

    # 画面: 項目定義の見出しが崩れていて数えられない。イベント定義は見出しだけで 0 件
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "画面項目定義"
    header(ws, [("画面ID", "S02"), ("画面名", "壊れた画面")])
    table(ws, 3, "画面項目定義", ["No", "項目"], [[1, "何か"]])
    ev = wb.create_sheet("イベント定義")
    table(ev, 3, "イベント定義", ["イベントID", "イベント名", "使用テーブル"], [])
    save(wb, "画面設計書_S02_壊れ.xlsx")

    # テーブル: M_USER はカラム 25（普通）、T_LOGIN_LOG はカラム 3（簡易）
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "M_USER"
    header(ws, [("テーブルID", "M_USER"), ("テーブル名", "ユーザーマスタ")])
    table(ws, 3, "カラム定義", ["No", "論理名", "物理名", "データ型"], [[i, "項目{0}".format(i), "COL_{0}".format(i), "CHAR"] for i in range(1, 26)])
    ws = wb.create_sheet("T_LOGIN_LOG")
    header(ws, [("テーブルID", "T_LOGIN_LOG"), ("テーブル名", "ログイン履歴")])
    table(ws, 3, "カラム定義", ["No", "論理名", "物理名", "データ型"], [[1, "ユーザーID", "USER_ID", "CHAR"], [2, "日時", "LOGIN_AT", "TIMESTAMP"], [3, "結果", "RESULT", "CHAR"]])
    save(wb, "テーブル定義書.xlsx")

    # IF: 受信（外部入力）。帳票: 外部出力
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "IF概要"
    header(ws, [("IF ID", "IF01"), ("IF名", "ユーザー連携")])
    ws["A3"] = "送受信区分"; ws["B3"] = "受信"
    rec = wb.create_sheet("レコード定義")
    table(rec, 3, "レコード定義", ["No", "項目名", "項目ID"], [[1, "ユーザーID", "USER_ID"], [2, "氏名", "NAME"]])
    save(wb, "IF定義書_IF01.xlsx")

    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "帳票項目"
    header(ws, [("帳票ID", "R01"), ("帳票名", "ログイン履歴一覧")])
    table(ws, 3, "帳票項目", ["No", "項目名"], [[i, "項目{0}".format(i)] for i in range(1, 5)])
    save(wb, "帳票設計書_R01.xlsx")


def main():
    tmp = tempfile.mkdtemp()
    try:
        make_grep_fixtures(tmp)
        make_pdf_fixtures()
        make_review_fixtures()
        make_deliverable_fixtures(tmp)
        make_view_fixtures()
        make_compare_fixtures()
        make_compare_format_fixtures(tmp)
        make_replace_fixtures(tmp)
        make_terms_fixtures()
        make_print_fixtures()
        make_size_fixtures()
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    print("fixtures written to", HERE)


if __name__ == "__main__":
    sys.exit(main())
