# -*- coding: utf-8 -*-
"""詳細設計サンプル（Excel 方眼紙）を生成する

    python make_samples.py

design_data.py の内容から、docs/design-sample/ 直下に設計書の xlsx を出力する。
"""
import math
import os
import sys
import unicodedata

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import design_data as D  # noqa: E402

OUT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# ---------------------------------------------------------------- 書式

GRID_COLS = 58          # 方眼の列数（A〜BF）
GRID_WIDTH = 2.75       # 方眼 1 マスの列幅
FONT_NAME = "Meiryo UI"

FONT = Font(name=FONT_NAME, size=9)
FONT_BOLD = Font(name=FONT_NAME, size=9, bold=True)
FONT_TITLE = Font(name=FONT_NAME, size=14, bold=True)
FONT_SECTION = Font(name=FONT_NAME, size=10, bold=True, color="1F4E79")
FONT_WHITE = Font(name=FONT_NAME, size=9, bold=True, color="FFFFFF")
FONT_MARK = Font(name=FONT_NAME, size=9, bold=True, color="C00000")

FILL_HEAD = PatternFill("solid", fgColor="DDEBF7")     # 表の見出し
FILL_LABEL = PatternFill("solid", fgColor="F2F2F2")    # ヘッダ部の項目名
FILL_BAR = PatternFill("solid", fgColor="1F4E79")      # 画面タイトルバー
FILL_BUTTON = PatternFill("solid", fgColor="D9D9D9")   # 画面のボタン
FILL_SCREEN = PatternFill("solid", fgColor="FFFFFF")

THIN = Side(style="thin", color="808080")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

LEFT = Alignment(horizontal="left", vertical="top", wrap_text=True)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
MIDDLE_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

MARKS = "①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬⑭⑮⑯⑰⑱⑲⑳"

# ---------------------------------------------------------------- 方眼紙のヘルパー


def new_sheet(wb, title):
    ws = wb.create_sheet(title)
    for c in range(1, GRID_COLS + 1):
        ws.column_dimensions[get_column_letter(c)].width = GRID_WIDTH
    ws.sheet_format.defaultRowHeight = 16
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    return ws


def put(ws, row, col, span, value, *, font=FONT, fill=None, align=LEFT, border=True, rows=1, fmt=None):
    """方眼の (row, col) から span マス分を結合して値を書く"""
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = font
    cell.alignment = align
    if fmt:
        cell.number_format = fmt
    last_col = col + span - 1
    last_row = row + rows - 1
    if span > 1 or rows > 1:
        ws.merge_cells(start_row=row, start_column=col, end_row=last_row, end_column=last_col)
    # 結合セルの罫線・塗りは範囲内の全セルに付けないと一部が欠ける
    for r in range(row, last_row + 1):
        for c in range(col, last_col + 1):
            target = ws.cell(row=r, column=c)
            if border:
                target.border = BORDER
            if fill:
                target.fill = fill
    return cell


def text_width(text):
    """全角を 1、半角を 0.6 として文字幅を数える（Meiryo UI の半角は全角の約 6 割）"""
    w = 0.0
    for ch in str(text):
        w += 1.0 if unicodedata.east_asian_width(ch) in ("F", "W", "A") else 0.6
    return w


def needed_lines(text, span):
    if text is None or text == "":
        return 1
    per_line = max(1.0, span * 1.7)   # 1 マスに全角約 1.7 文字（余白を見込む）
    return sum(max(1, math.ceil(text_width(line) / per_line)) for line in str(text).split("\n"))


def set_height(ws, row, lines):
    ws.row_dimensions[row].height = max(16, lines * 13 + 4)


def section(ws, row, title):
    put(ws, row, 1, GRID_COLS, "■ " + title, font=FONT_SECTION, border=False, align=MIDDLE_LEFT)
    ws.row_dimensions[row].height = 20
    return row + 1


def table(ws, top, columns, rows, *, col_start=1):
    """columns: [(見出し, マス数)]。見出し行と明細行を書き、次の行番号を返す"""
    c = col_start
    for label, span in columns:
        put(ws, top, c, span, label, font=FONT_BOLD, fill=FILL_HEAD, align=CENTER)
        c += span
    ws.row_dimensions[top].height = 20

    r = top + 1
    for values in rows:
        c = col_start
        lines = 1
        for (label, span), value in zip(columns, values):
            fmt = "yyyy/mm/dd" if hasattr(value, "year") else None
            align = CENTER if span <= 3 else LEFT
            put(ws, r, c, span, value, align=align, fmt=fmt)
            lines = max(lines, needed_lines(value if not hasattr(value, "year") else "", span))
            c += span
        set_height(ws, r, lines)
        r += 1
    return r


# 作成中のブックの最終更新（改訂履歴の最後の行の改訂日・改訂者）。表紙と各シートのヘッダに使う
_last_update = {"date": D.CREATED, "by": D.AUTHOR}


def doc_header(ws, title, id_label="", id_value="", name_label="", name_value=""):
    """各シート共通の帳票ヘッダ（1〜2 行目）"""
    put(ws, 1, 1, 14, title, font=FONT_TITLE, align=CENTER, rows=2)
    put(ws, 1, 15, 4, "システム名", font=FONT_BOLD, fill=FILL_LABEL, align=CENTER)
    put(ws, 1, 19, 16, D.SYSTEM, align=MIDDLE_LEFT)
    put(ws, 1, 35, 4, "作成者", font=FONT_BOLD, fill=FILL_LABEL, align=CENTER)
    put(ws, 1, 39, 8, D.AUTHOR, align=MIDDLE_LEFT)
    put(ws, 1, 47, 4, "作成日", font=FONT_BOLD, fill=FILL_LABEL, align=CENTER)
    put(ws, 1, 51, 8, D.CREATED, align=MIDDLE_LEFT, fmt="yyyy/mm/dd")
    put(ws, 2, 15, 4, id_label or "サブシステム", font=FONT_BOLD, fill=FILL_LABEL, align=CENTER)
    if name_label:
        put(ws, 2, 19, 6, id_value, align=MIDDLE_LEFT)
        put(ws, 2, 25, 4, name_label, font=FONT_BOLD, fill=FILL_LABEL, align=CENTER)
        put(ws, 2, 29, 6, name_value, align=MIDDLE_LEFT)
    else:
        put(ws, 2, 19, 16, id_value or D.SUBSYSTEM, align=MIDDLE_LEFT)
    put(ws, 2, 35, 4, "更新者", font=FONT_BOLD, fill=FILL_LABEL, align=CENTER)
    put(ws, 2, 39, 8, _last_update["by"], align=MIDDLE_LEFT)
    put(ws, 2, 47, 4, "更新日", font=FONT_BOLD, fill=FILL_LABEL, align=CENTER)
    put(ws, 2, 51, 8, _last_update["date"], align=MIDDLE_LEFT, fmt="yyyy/mm/dd")
    ws.row_dimensions[1].height = 18
    ws.row_dimensions[2].height = 18
    return 4


# ---------------------------------------------------------------- 表紙・改訂履歴


def cover_sheet(wb, doc_title, doc_id, version):
    ws = new_sheet(wb, "表紙")
    put(ws, 10, 1, GRID_COLS, D.SYSTEM, font=Font(name=FONT_NAME, size=16, bold=True), align=CENTER, border=False)
    put(ws, 12, 1, GRID_COLS, doc_title, font=Font(name=FONT_NAME, size=22, bold=True), align=CENTER, border=False)
    ws.row_dimensions[12].height = 36
    put(ws, 14, 1, GRID_COLS, "詳細設計書", font=Font(name=FONT_NAME, size=12), align=CENTER, border=False)

    rows = [("文書番号", doc_id), ("版数", version), ("作成日", D.CREATED), ("最終更新日", _last_update["date"]),
            ("作成者", D.AUTHOR), ("作成部署", D.COMPANY)]
    r = 22
    for label, value in rows:
        put(ws, r, 17, 7, label, font=FONT_BOLD, fill=FILL_LABEL, align=CENTER)
        put(ws, r, 24, 18, value, align=MIDDLE_LEFT, fmt="yyyy/mm/dd" if hasattr(value, "year") else None)
        ws.row_dimensions[r].height = 20
        r += 1
    put(ws, 34, 1, GRID_COLS, "※ 本書は設計ツール検証用のサンプルです。システム・会社・人物はすべて架空です。",
        font=Font(name=FONT_NAME, size=9, color="808080"), align=CENTER, border=False)


def history_sheet(wb, doc_title, history):
    ws = new_sheet(wb, "改訂履歴")
    r = doc_header(ws, doc_title)
    r = section(ws, r, "改訂履歴")
    table(ws, r, [("版数", 4), ("改訂日", 6), ("改訂箇所", 12), ("改訂内容", 28), ("改訂者", 8)],
          [list(h) for h in history])


def new_book(doc_title, doc_id, version, history):
    # 表紙の版数・最終更新日と、改訂履歴の最後の行をそろえる（Test-ExcelDeliverable.ps1 の版数チェックで矛盾しないように）
    assert history[-1][0] == version, (doc_title, version, history[-1])
    _last_update["date"], _last_update["by"] = history[-1][1], history[-1][4]
    wb = Workbook()
    wb.remove(wb.active)
    cover_sheet(wb, doc_title, doc_id, version)
    history_sheet(wb, doc_title, history)
    return wb


def save(wb, filename):
    # 開いたときに先頭シートの A1 が表示されるようにしておく
    wb.active = 0
    for ws in wb.worksheets:
        ws.sheet_view.selection[0].activeCell = "A1"
        ws.sheet_view.selection[0].sqref = "A1"
        ws.sheet_view.tabSelected = ws is wb.worksheets[0]
    # 作成者のプロパティ（既定は「openpyxl」）は残さない（提出前チェックで指摘されないように）
    wb.properties.creator = None
    path = os.path.join(OUT_DIR, filename)
    wb.save(path)
    print("wrote", os.path.relpath(path, OUT_DIR))


# ---------------------------------------------------------------- 画面一覧


def screen_name(screen_id):
    for s in D.SCREENS:
        if s["id"] == screen_id:
            return s["name"]
    return ""


def make_screen_list():
    title = "画面一覧"
    wb = new_book(title, "ODR-DD-01", "1.0", D.DEFAULT_HISTORY)

    ws = new_sheet(wb, "画面一覧")
    r = doc_header(ws, title)
    r = section(ws, r, "画面一覧")
    table(ws, r, [("No", 2), ("画面ID", 5), ("画面名", 8), ("画面区分", 6), ("概要", 25), ("利用者", 6), ("設計書", 6)],
          [(i + 1, s["id"], s["name"], s["kind"], s["summary"], s["users"], "DD02") for i, s in enumerate(D.SCREENS)])

    ws = new_sheet(wb, "画面遷移")
    r = doc_header(ws, title)
    r = section(ws, r, "画面遷移")
    table(ws, r, [("No", 2), ("遷移元画面ID", 6), ("遷移元画面名", 8), ("イベント", 9), ("遷移先画面ID", 6),
                  ("遷移先画面名", 8), ("引き渡し項目", 8), ("備考", 11)],
          [(i + 1, src, screen_name(src), ev, dst, screen_name(dst), param, note)
           for i, (src, ev, dst, param, note) in enumerate(D.TRANSITIONS)])

    save(wb, "DD01_画面一覧.xlsx")


# ---------------------------------------------------------------- 画面設計書


def item_by_no(screen, no):
    for it in screen["items"]:
        if it[0] == no:
            return it
    raise KeyError(no)


def mark(no):
    return MARKS[no - 1] if no <= len(MARKS) else "({})".format(no)


def draw_frame(ws, top, bottom, title):
    """画面の外枠とタイトルバーを描く。枠の内側に描いた入力欄の罫線は残す"""
    left, right = 2, GRID_COLS - 1
    for r in range(top, bottom + 1):
        for c in range(left, right + 1):
            if left < c < right and top < r < bottom:
                continue
            cell = ws.cell(row=r, column=c)
            old = cell.border
            cell.border = Border(
                left=THIN if c == left else old.left, right=THIN if c == right else old.right,
                top=THIN if r == top else old.top, bottom=THIN if r == bottom else old.bottom)
    put(ws, top, left, right - left + 1, "  " + title, font=FONT_WHITE, fill=FILL_BAR, align=MIDDLE_LEFT)
    ws.row_dimensions[top].height = 20


def field(ws, row, col, screen, no, label_span, box_span, sample=""):
    """項目名ラベル＋入力欄。項目名の前に項目定義の No を丸数字で付ける"""
    it = item_by_no(screen, no)
    put(ws, row, col, 1, mark(no), font=FONT_MARK, align=CENTER, border=False)
    put(ws, row, col + 1, label_span, it[1], font=FONT_BOLD, align=MIDDLE_LEFT, border=False)
    put(ws, row, col + 1 + label_span, box_span, sample, align=MIDDLE_LEFT)


def button(ws, row, col, screen, no, span=5):
    it = item_by_no(screen, no)
    put(ws, row, col, 1, mark(no), font=FONT_MARK, align=CENTER, border=False)
    put(ws, row, col + 1, span, it[1], font=FONT_BOLD, fill=FILL_BUTTON, align=CENTER)


def grid_list(ws, top, col, screen, list_no, columns, rows):
    """一覧（明細表）を描く。columns: [(項目No, マス数)]"""
    put(ws, top - 1, col, 1, mark(list_no), font=FONT_MARK, align=CENTER, border=False)
    put(ws, top - 1, col + 1, 10, item_by_no(screen, list_no)[1], font=FONT_BOLD, align=MIDDLE_LEFT, border=False)
    c = col
    for no, span in columns:
        put(ws, top, c, span, mark(no) + item_by_no(screen, no)[1], font=FONT_BOLD, fill=FILL_HEAD, align=CENTER)
        c += span
    r = top + 1
    for values in rows:
        c = col
        for (no, span), v in zip(columns, values):
            put(ws, r, c, span, v, align=CENTER if span <= 4 else MIDDLE_LEFT)
            c += span
        r += 1
    return r


def layout_search(ws, s, top):
    r = top + 2
    field(ws, r, 4, s, 1, 6, 9, "2026")
    field(ws, r, 22, s, 2, 7, 7, "2026/09/15")
    put(ws, r, 37, 2, "～", align=CENTER, border=False)
    field(ws, r, 39, s, 3, 6, 7, "2026/10/15")
    r += 2
    field(ws, r, 4, s, 4, 6, 9)
    field(ws, r, 22, s, 5, 7, 18, "サンプル")
    r += 2
    put(ws, r, 4, 1, mark(6), font=FONT_MARK, align=CENTER, border=False)
    put(ws, r, 5, 6, item_by_no(s, 6)[1], font=FONT_BOLD, align=MIDDLE_LEFT, border=False)
    put(ws, r, 11, 30, "☑受付　☐出荷指示済　☐キャンセル", align=MIDDLE_LEFT, border=False)
    r += 2
    button(ws, r, 34, s, 7)
    button(ws, r, 41, s, 8)
    button(ws, r, 48, s, 9)
    r += 2
    put(ws, r, 4, 1, mark(10), font=FONT_MARK, align=CENTER, border=False)
    put(ws, r, 5, 8, "検索結果：3件", align=MIDDLE_LEFT, border=False)
    r += 2
    r = grid_list(ws, r, 4, s, 11, [(12, 8), (13, 7), (14, 14), (15, 8), (16, 8), (17, 7)], [
        ("202610150002", "2026/10/15", "鈴木 一郎", "受付", "12,800", ""),
        ("202610150001", "2026/10/15", "株式会社サンプル商事", "受付", "45,000", "2026/10/20"),
        ("202610140007", "2026/10/14", "有限会社テスト工業", "出荷指示済", "8,500", "2026/10/16"),
    ])
    put(ws, r, 4, 50, "＜ 前へ　1　次へ ＞", align=CENTER, border=False)
    return r + 1


def layout_entry(ws, s, top):
    r = top + 2
    field(ws, r, 4, s, 1, 6, 8, "2026/10/15")
    r += 2
    field(ws, r, 4, s, 2, 6, 8, "C0000123")
    field(ws, r, 22, s, 3, 5, 20, "株式会社サンプル商事")
    r += 2
    field(ws, r, 4, s, 4, 6, 8, "銀行振込 ▼")
    field(ws, r, 22, s, 5, 5, 8, "2026/10/20")
    field(ws, r, 39, s, 6, 6, 8, "午前中 ▼")
    r += 2
    field(ws, r, 4, s, 7, 6, 46, "")
    r += 2
    r = grid_list(ws, r, 4, s, 8, [(9, 4), (10, 8), (11, 14), (12, 7), (13, 6), (14, 7), (15, 4)], [
        ("1", "P000000101", "業務用ボールペン（黒）10本入", "1,000", "3", "3,000", "☐"),
        ("2", "P000000205", "A4コピー用紙 2,500枚", "42,000", "1", "42,000", "☐"),
    ])
    button(ws, r, 4, s, 16)
    put(ws, r, 34, 1, mark(17), font=FONT_MARK, align=CENTER, border=False)
    put(ws, r, 35, 6, item_by_no(s, 17)[1], font=FONT_BOLD, align=MIDDLE_LEFT, border=False)
    put(ws, r, 41, 9, "45,000", align=MIDDLE_LEFT)
    r += 2
    button(ws, r, 40, s, 18)
    button(ws, r, 47, s, 19)
    return r + 1


def layout_detail(ws, s, top):
    r = top + 2
    field(ws, r, 4, s, 1, 6, 9, "202610150001")
    field(ws, r, 22, s, 2, 5, 8, "2026/10/15")
    field(ws, r, 38, s, 3, 7, 7, "受付")
    r += 2
    field(ws, r, 4, s, 4, 6, 9, "C0000123")
    field(ws, r, 22, s, 5, 5, 16, "株式会社サンプル商事")
    r += 2
    field(ws, r, 4, s, 6, 6, 9, "銀行振込")
    field(ws, r, 22, s, 7, 5, 8, "2026/10/20")
    field(ws, r, 38, s, 8, 7, 7, "午前中 ▼")
    r += 2
    field(ws, r, 4, s, 9, 6, 46, "")
    r += 2
    r = grid_list(ws, r, 4, s, 10, [(11, 4), (12, 8), (13, 16), (14, 7), (15, 6), (16, 9)], [
        ("1", "P000000101", "業務用ボールペン（黒）10本入", "1,000", "3", "3,000"),
        ("2", "P000000205", "A4コピー用紙 2,500枚", "42,000", "1", "42,000"),
    ])
    put(ws, r, 34, 1, mark(17), font=FONT_MARK, align=CENTER, border=False)
    put(ws, r, 35, 6, item_by_no(s, 17)[1], font=FONT_BOLD, align=MIDDLE_LEFT, border=False)
    put(ws, r, 41, 9, "45,000", align=MIDDLE_LEFT)
    r += 2
    button(ws, r, 33, s, 18)
    button(ws, r, 40, s, 19)
    button(ws, r, 47, s, 20)
    return r + 1


LAYOUTS = {"search": layout_search, "entry": layout_entry, "detail": layout_detail}


def make_screen(s, index):
    title = "画面設計書"
    doc_id = "ODR-DD-02-{0:02d}".format(index)
    wb = new_book("{0}（{1} {2}）".format(title, s["id"], s["name"]), doc_id, "1.0", D.DEFAULT_HISTORY)

    # 画面概要
    ws = new_sheet(wb, "画面概要")
    r = doc_header(ws, title, "画面ID", s["id"], "画面名", s["name"])
    r = section(ws, r, "画面概要")
    for label, value in [("画面ID", s["id"]), ("画面名", s["name"]), ("画面区分", s["kind"]),
                         ("概要", s["summary"]), ("利用者", s["users"])]:
        put(ws, r, 1, 8, label, font=FONT_BOLD, fill=FILL_LABEL, align=CENTER)
        put(ws, r, 9, 50, value, align=MIDDLE_LEFT)
        set_height(ws, r, needed_lines(value, 50))
        r += 1
    r += 1
    r = section(ws, r, "遷移先")
    rows = [(i + 1, ev, dst, screen_name(dst), param, note)
            for i, (src, ev, dst, param, note) in enumerate([t for t in D.TRANSITIONS if t[0] == s["id"]])]
    table(ws, r, [("No", 2), ("イベント", 10), ("遷移先画面ID", 6), ("遷移先画面名", 10), ("引き渡し項目", 10), ("備考", 20)], rows)

    # 画面レイアウト
    ws = new_sheet(wb, "画面レイアウト")
    r = doc_header(ws, title, "画面ID", s["id"], "画面名", s["name"])
    r = section(ws, r, "画面レイアウト（丸数字は画面項目定義の No）")
    top = r + 1
    bottom = LAYOUTS[s["layout"]](ws, s, top) + 1
    draw_frame(ws, top, bottom, "{0}　　{1}".format(D.SYSTEM, s["name"]))
    hidden = [it for it in s["items"] if it[3] == "隠し項目"]
    if hidden:
        put(ws, bottom + 2, 2, 50, "※ 隠し項目（画面には表示しない）：" + "、".join(
            "{0}{1}".format(mark(it[0]), it[1]) for it in hidden), border=False)

    # 画面項目定義
    ws = new_sheet(wb, "画面項目定義")
    r = doc_header(ws, title, "画面ID", s["id"], "画面名", s["name"])
    r = section(ws, r, "画面項目定義")
    table(ws, r, [("No", 2), ("項目名", 6), ("項目ID", 7), ("区分", 5), ("入出力", 3), ("型", 4), ("桁数", 2),
                  ("必須", 2), ("初期値", 5), ("入力チェック", 8), ("対応項目", 8), ("備考", 6)],
          [it for it in s["items"]])

    # イベント定義
    ws = new_sheet(wb, "イベント定義")
    r = doc_header(ws, title, "画面ID", s["id"], "画面名", s["name"])
    r = section(ws, r, "イベント定義（使用テーブルの C:登録 R:参照 U:更新 D:削除）")
    table(ws, r, [("イベントID", 4), ("イベント名", 5), ("トリガー", 6), ("処理概要", 23), ("使用テーブル", 8),
                  ("メッセージ", 7), ("遷移先", 5)],
          [ev for ev in s["events"]])

    save(wb, "DD02_画面設計書_{0}_{1}.xlsx".format(s["id"], s["name"]))


# ---------------------------------------------------------------- テーブル定義書


def make_tables():
    title = "テーブル定義書"
    wb = new_book(title, "ODR-DD-03", "1.1", D.TABLE_HISTORY)

    ws = new_sheet(wb, "テーブル一覧")
    r = doc_header(ws, title)
    r = section(ws, r, "テーブル一覧")
    table(ws, r, [("No", 2), ("テーブルID", 9), ("テーブル名", 9), ("区分", 7), ("概要", 25), ("想定件数", 6)],
          [(i + 1, t[0], t[1], t[2], t[3], t[4]) for i, t in enumerate(D.TABLES)])

    for tid, tname, kind, summary, volume, columns, indexes in D.TABLES:
        ws = new_sheet(wb, tid)
        r = doc_header(ws, title, "テーブルID", tid, "テーブル名", tname)
        r = section(ws, r, "テーブル概要")
        for label, value in [("テーブルID", tid), ("テーブル名", tname), ("区分", kind), ("概要", summary), ("想定件数", volume)]:
            put(ws, r, 1, 8, label, font=FONT_BOLD, fill=FILL_LABEL, align=CENTER)
            put(ws, r, 9, 50, value, align=MIDDLE_LEFT)
            r += 1
        r += 1
        r = section(ws, r, "カラム定義")
        rows = []
        for i, (lname, pname, dtype, length, scale, pk, notnull, default, fk, note) in enumerate(columns):
            rows.append((i + 1, lname, pname, dtype, length, scale, pk, "○" if notnull else "", default, fk, note))
        r = table(ws, r, [("No", 2), ("論理名", 7), ("物理名", 9), ("データ型", 5), ("桁数", 3), ("小数", 2),
                          ("PK", 2), ("必須", 2), ("デフォルト", 8), ("外部キー参照先", 9), ("備考", 9)], rows)
        r += 1
        r = section(ws, r, "インデックス定義")
        if indexes:
            table(ws, r, [("No", 2), ("インデックス名", 12), ("カラム", 18), ("ユニーク", 4), ("備考", 22)],
                  [(i + 1, name, cols, "○" if uniq else "", "") for i, (name, cols, uniq) in enumerate(indexes)])
        else:
            put(ws, r, 1, 58, "なし（主キーのみ）", border=False)

    save(wb, "DD03_テーブル定義書.xlsx")


# ---------------------------------------------------------------- コード定義書・メッセージ一覧


def make_codes():
    title = "コード定義書"
    wb = new_book(title, "ODR-DD-04", "1.0", D.DEFAULT_HISTORY)
    ws = new_sheet(wb, "コード一覧")
    r = doc_header(ws, title)
    r = section(ws, r, "コード一覧（コード値はアプリケーションの定数として保持する）")
    rows = []
    for cid, cname, desc, values in D.CODES:
        for i, (value, label, note) in enumerate(values):
            rows.append((cid, cname, value, label, i + 1, note))
    table(ws, r, [("コード区分ID", 6), ("コード区分名", 9), ("コード値", 5), ("表示名", 10), ("表示順", 4), ("備考", 24)], rows)
    save(wb, "DD04_コード定義書.xlsx")


def make_messages():
    title = "メッセージ一覧"
    wb = new_book(title, "ODR-DD-05", "1.0", D.DEFAULT_HISTORY)
    ws = new_sheet(wb, "メッセージ一覧")
    r = doc_header(ws, title)
    r = section(ws, r, "メッセージ一覧（{n} は実行時に置き換えるパラメータ）")
    table(ws, r, [("メッセージID", 6), ("種別", 4), ("メッセージ", 30), ("パラメータ", 18)], D.MESSAGES)
    save(wb, "DD05_メッセージ一覧.xlsx")


# ---------------------------------------------------------------- IF 定義書


def make_interface(itf, index):
    title = "IF定義書"
    wb = new_book("{0}（{1} {2}）".format(title, itf["id"], itf["name"]), "ODR-DD-06-{0:02d}".format(index),
                  "1.0", D.DEFAULT_HISTORY)

    ws = new_sheet(wb, "IF概要")
    r = doc_header(ws, title, "IF ID", itf["id"], "IF名", itf["name"])
    r = section(ws, r, "IF概要")
    for label, value in itf["overview"]:
        put(ws, r, 1, 10, label, font=FONT_BOLD, fill=FILL_LABEL, align=CENTER)
        put(ws, r, 11, 48, value, align=MIDDLE_LEFT)
        set_height(ws, r, needed_lines(value, 48))
        r += 1

    ws = new_sheet(wb, "レコード定義")
    r = doc_header(ws, title, "IF ID", itf["id"], "IF名", itf["name"])
    r = section(ws, r, "レコード定義")
    r = table(ws, r, [("No", 2), ("項目名", 7), ("項目ID", 9), ("属性", 3), ("桁数", 3), ("必須", 2),
                      ("編集形式", 6), ("設定元", 11), ("備考", 15)], itf["records"])
    r += 1
    r = section(ws, r, "データ例")
    for line in itf["sample"]:
        put(ws, r, 1, 58, line, font=Font(name="Consolas", size=8), align=MIDDLE_LEFT)
        r += 1

    save(wb, "DD06_IF定義書_{0}_{1}.xlsx".format(itf["id"], itf["name"]))


def main():
    make_screen_list()
    for i, s in enumerate(D.SCREENS, start=1):
        make_screen(s, i)
    make_tables()
    make_codes()
    make_messages()
    for i, itf in enumerate(D.INTERFACES, start=1):
        make_interface(itf, i)


if __name__ == "__main__":
    main()
